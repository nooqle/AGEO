"""Services for understanding uploaded CSV/XLSX tables."""

from __future__ import annotations

import csv
import json
import logging
import os
from dataclasses import dataclass
from io import StringIO
from typing import Any
from uuid import UUID

from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.llm import get_llm_model
from app.core.utils import extract_json_from_content
from app.models.file_metadata import FileMetadata

logger = logging.getLogger(__name__)

_QUESTION_HEADER_KEYWORDS = (
    "问题",
    "question",
    "query",
    "ask",
    "提问",
)
_CATEGORY_HEADER_KEYWORDS = ("分类", "category", "topic", "主题", "场景")
_BRAND_HEADER_KEYWORDS = ("品牌", "brand", "官网", "website", "domain", "关键词")
_COMPETITOR_HEADER_KEYWORDS = ("竞品", "competitor", "对手")
_LINK_HEADER_KEYWORDS = ("链接", "link", "url", "网址", "来源", "source", "domain")
_INDUSTRY_HEADER_KEYWORDS = ("行业", "industry", "赛道", "品类")
_KEYWORD_HEADER_KEYWORDS = ("关键词", "keyword", "核心词")
_PRODUCT_HEADER_KEYWORDS = ("产品", "product", "品类", "产品线")


@dataclass
class ParsedTable:
    file_name: str
    file_id: str
    mime_type: str
    sheet_name: str | None
    headers: list[str]
    rows: list[dict[str, str]]
    raw_row_count: int
    valid_row_count: int
    duplicate_row_count: int
    empty_row_count: int


class TableIntakeService:
    """Reads uploaded table files and classifies them into business shapes."""

    def __init__(self, db: AsyncSession):
        self.db = db

    async def analyze_attachment(
        self,
        *,
        attachment: dict[str, Any],
        user_message: str = "",
        brand_context: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        file_id = str(attachment.get("file_id") or attachment.get("id") or "").strip()
        if not file_id:
            raise ValueError("附件缺少 file_id")

        file_meta = await self.db.get(FileMetadata, UUID(file_id))
        if file_meta is None:
            raise ValueError("未找到上传文件")

        parsed = self._parse_file(file_meta)
        deterministic = self._deterministic_classify(parsed)
        llm_result = await self._llm_classify_if_needed(
            parsed=parsed,
            deterministic=deterministic,
            user_message=user_message,
            brand_context=brand_context or {},
        )
        result = self._merge_results(parsed, deterministic, llm_result)
        return result

    def _parse_file(self, file_meta: FileMetadata) -> ParsedTable:
        ext = os.path.splitext(file_meta.name or "")[1].lower()
        if ext == ".csv":
            return self._parse_csv(file_meta)
        if ext == ".xlsx":
            return self._parse_xlsx(file_meta)
        raise ValueError("仅支持 CSV 或 XLSX 表格")

    def _parse_csv(self, file_meta: FileMetadata) -> ParsedTable:
        with open(file_meta.path, "rb") as f:
            raw = f.read()

        text = raw.decode("utf-8-sig", errors="ignore")
        reader = csv.reader(StringIO(text))
        rows = list(reader)
        return self._build_parsed_table(
            file_name=file_meta.name,
            file_id=str(file_meta.id),
            mime_type=file_meta.content_type,
            sheet_name=None,
            matrix=rows,
        )

    def _parse_xlsx(self, file_meta: FileMetadata) -> ParsedTable:
        workbook = load_workbook(file_meta.path, read_only=True, data_only=True)
        for worksheet in workbook.worksheets:
            matrix: list[list[str]] = []
            for row in worksheet.iter_rows(values_only=True):
                matrix.append(
                    ["" if cell is None else str(cell).strip() for cell in row]
                )
            if any(any(cell for cell in row) for row in matrix):
                return self._build_parsed_table(
                    file_name=file_meta.name,
                    file_id=str(file_meta.id),
                    mime_type=file_meta.content_type,
                    sheet_name=worksheet.title,
                    matrix=matrix,
                )
        raise ValueError("表格为空，未解析到有效数据")

    def _build_parsed_table(
        self,
        *,
        file_name: str,
        file_id: str,
        mime_type: str,
        sheet_name: str | None,
        matrix: list[list[str]],
    ) -> ParsedTable:
        rows_after_header = matrix[1:] if len(matrix) > 1 else []
        empty_row_count = sum(
            1 for row in rows_after_header if not any(str(cell).strip() for cell in row)
        )
        non_empty_rows = [row for row in matrix if any(str(cell).strip() for cell in row)]
        if len(non_empty_rows) < 2:
            raise ValueError("表格内容不足，至少需要表头和一行数据")

        header_row = [str(cell).strip() for cell in non_empty_rows[0]]
        if not any(header_row):
            raise ValueError("表头为空，无法解析")

        seen_headers: dict[str, int] = {}
        headers: list[str] = []
        for index, header in enumerate(header_row, start=1):
            base = header or f"column_{index}"
            count = seen_headers.get(base, 0) + 1
            seen_headers[base] = count
            headers.append(base if count == 1 else f"{base}_{count}")

        rows: list[dict[str, str]] = []
        seen_signatures: set[str] = set()
        duplicate_row_count = 0

        for raw_row in non_empty_rows[1:]:
            values = [str(cell).strip() for cell in raw_row]
            padded = values + [""] * max(0, len(headers) - len(values))
            row_dict = {headers[i]: padded[i] if i < len(padded) else "" for i in range(len(headers))}
            signature = json.dumps(row_dict, ensure_ascii=False, sort_keys=True)
            if signature in seen_signatures:
                duplicate_row_count += 1
                continue
            seen_signatures.add(signature)
            rows.append(row_dict)

        if not rows:
            raise ValueError("表格没有可用数据行")

        return ParsedTable(
            file_name=file_name,
            file_id=file_id,
            mime_type=mime_type,
            sheet_name=sheet_name,
            headers=headers,
            rows=rows,
            raw_row_count=max(len(matrix) - 1, 0),
            valid_row_count=len(rows),
            duplicate_row_count=duplicate_row_count,
            empty_row_count=empty_row_count,
        )

    def _deterministic_classify(self, parsed: ParsedTable) -> dict[str, Any]:
        header_map = {header.lower(): header for header in parsed.headers}
        question_header = self._match_header(parsed.headers, _QUESTION_HEADER_KEYWORDS)
        category_header = self._match_header(parsed.headers, _CATEGORY_HEADER_KEYWORDS)
        link_header = self._match_header(parsed.headers, _LINK_HEADER_KEYWORDS)
        brand_related_headers = [
            header
            for header in parsed.headers
            if self._contains_any(header, _BRAND_HEADER_KEYWORDS + _COMPETITOR_HEADER_KEYWORDS)
        ]

        if question_header:
            questions, skipped_count = self._normalize_question_rows(
                parsed.rows, question_header, category_header
            )
            return {
                "table_kind": "question_list",
                "confidence": 0.96,
                "recommended_step": "A3",
                "detected_columns": {
                    "question": question_header,
                    **({"category": category_header} if category_header else {}),
                },
                "normalized_payload": {
                    "questions": questions,
                },
                "warnings": self._build_warnings(parsed, skipped_question_count=skipped_count),
            }

        if link_header or self._looks_like_link_sheet(parsed):
            return {
                "table_kind": "link_list",
                "confidence": 0.88,
                "recommended_step": "CONFIDENCE_EVAL",
                "detected_columns": {"link": link_header} if link_header else {},
                "normalized_payload": {
                    "links": self._normalize_link_rows(parsed.rows, link_header),
                },
                "warnings": self._build_warnings(parsed),
            }

        if brand_related_headers or self._looks_like_brand_sheet(parsed):
            return {
                "table_kind": "brand_competitor_info",
                "confidence": 0.82,
                "recommended_step": "A1",
                "detected_columns": {header.lower(): header for header in brand_related_headers[:6]},
                "normalized_payload": {
                    **self._normalize_brand_rows(parsed.rows, parsed.headers),
                    "rows": parsed.rows,
                },
                "warnings": self._build_warnings(parsed),
            }

        single_col_header = parsed.headers[0] if len(parsed.headers) == 1 else None
        if single_col_header and self._looks_like_freeform_question_rows(parsed.rows, single_col_header):
            questions, skipped_count = self._normalize_question_rows(
                parsed.rows, single_col_header, None
            )
            return {
                "table_kind": "question_list",
                "confidence": 0.78,
                "recommended_step": "A3",
                "detected_columns": {"question": single_col_header},
                "normalized_payload": {"questions": questions},
                "warnings": self._build_warnings(parsed, skipped_question_count=skipped_count),
            }

        return {
            "table_kind": "unknown",
            "confidence": 0.4,
            "recommended_step": "UNKNOWN",
            "detected_columns": header_map,
            "normalized_payload": {"rows": parsed.rows[:20]},
            "warnings": self._build_warnings(parsed),
        }

    async def _llm_classify_if_needed(
        self,
        *,
        parsed: ParsedTable,
        deterministic: dict[str, Any],
        user_message: str,
        brand_context: dict[str, Any],
    ) -> dict[str, Any] | None:
        if deterministic.get("confidence", 0) >= 0.9:
            return None

        model = get_llm_model()
        prompt = {
            "file_name": parsed.file_name,
            "sheet_name": parsed.sheet_name,
            "headers": parsed.headers,
            "sample_rows": parsed.rows[:5],
            "user_message": user_message,
            "brand_context": {
                "brand_name": brand_context.get("brand_name"),
                "industry": brand_context.get("industry"),
            },
            "deterministic_guess": {
                "table_kind": deterministic.get("table_kind"),
                "confidence": deterministic.get("confidence"),
            },
        }

        system = (
            "你是表格导入理解器。请判断上传表格属于哪一类："
            "question_list、brand_competitor_info、link_list、unknown。"
            "只返回 JSON。"
        )
        user = (
            "请结合表头、样例数据、用户消息和品牌上下文，输出："
            '{"table_kind":"...", "confidence":0.0, "detected_columns":{"question":"问题列名"}, '
            '"reason":"一句话原因"}\n\n'
            f"{json.dumps(prompt, ensure_ascii=False)}"
        )

        try:
            response = await model.async_call(
                messages=[
                    {"role": "system", "content": system},
                    {"role": "user", "content": user},
                ],
                temperature=0.1,
                max_tokens=600,
            )
            data = extract_json_from_content(response.content if hasattr(response, "content") else str(response))
            if not isinstance(data, dict):
                return None
            return data
        except Exception as exc:
            logger.warning("[TableIntake] LLM classify failed: %s", exc)
            return None

    def _merge_results(
        self,
        parsed: ParsedTable,
        deterministic: dict[str, Any],
        llm_result: dict[str, Any] | None,
    ) -> dict[str, Any]:
        final = dict(deterministic)
        if llm_result:
            llm_kind = str(llm_result.get("table_kind") or "").strip()
            llm_confidence = float(llm_result.get("confidence") or 0)
            if llm_kind in {"question_list", "brand_competitor_info", "link_list", "unknown"} and llm_confidence >= final.get("confidence", 0):
                final["table_kind"] = llm_kind
                final["confidence"] = llm_confidence
                detected_columns = llm_result.get("detected_columns")
                if isinstance(detected_columns, dict) and detected_columns:
                    final["detected_columns"] = {
                        str(key): str(value)
                        for key, value in detected_columns.items()
                    }

        final["recommended_step"] = {
            "question_list": "A3",
            "brand_competitor_info": "A1",
            "link_list": "CONFIDENCE_EVAL",
        }.get(final.get("table_kind"), "UNKNOWN")
        final["needs_user_confirmation"] = True
        final["source_file"] = {
            "file_id": parsed.file_id,
            "name": parsed.file_name,
            "mime_type": parsed.mime_type,
            "sheet_name": parsed.sheet_name,
        }
        final["stats"] = {
            "row_count": parsed.raw_row_count,
            "valid_row_count": parsed.valid_row_count,
            "duplicate_row_count": parsed.duplicate_row_count,
            "empty_row_count": parsed.empty_row_count,
        }

        if final.get("table_kind") == "question_list":
            questions = list((final.get("normalized_payload") or {}).get("questions") or [])
            final["summary"] = f"识别为问题列表，共 {len(questions)} 条有效问题。"
        elif final.get("table_kind") == "brand_competitor_info":
            final["summary"] = f"识别为品牌/竞品信息表，共 {parsed.valid_row_count} 行。"
        elif final.get("table_kind") == "link_list":
            links = list((final.get("normalized_payload") or {}).get("links") or [])
            final["summary"] = f"识别为链接清单，共 {len(links)} 条。"
        else:
            final["summary"] = f"未能稳定识别表格用途，共读取 {parsed.valid_row_count} 行。"

        return final

    def _normalize_question_rows(
        self,
        rows: list[dict[str, str]],
        question_header: str,
        category_header: str | None,
    ) -> tuple[list[dict[str, Any]], int]:
        normalized: list[dict[str, Any]] = []
        skipped_count = 0
        for index, row in enumerate(rows, start=1):
            text = str(row.get(question_header, "")).strip()
            if not text:
                skipped_count += 1
                continue
            normalized.append(
                {
                    "id": f"upload_q_{index:03d}",
                    "text": text,
                    "category": str(row.get(category_header, "")).strip() if category_header else "",
                    "intent": "",
                    "stage": "",
                    "source": "uploaded_table",
                }
            )
        return normalized, skipped_count

    def _normalize_link_rows(
        self,
        rows: list[dict[str, str]],
        link_header: str | None,
    ) -> list[dict[str, Any]]:
        normalized: list[dict[str, Any]] = []
        for index, row in enumerate(rows, start=1):
            candidate = ""
            if link_header:
                candidate = str(row.get(link_header, "")).strip()
            else:
                candidate = next(
                    (str(value).strip() for value in row.values() if self._looks_like_url(str(value).strip())),
                    "",
                )
            if not candidate:
                continue
            normalized.append(
                {
                    "id": f"link_{index:03d}",
                    "url": candidate,
                    "label": row.get("名称") or row.get("name") or row.get("备注") or "",
                }
            )
        return normalized

    def _normalize_brand_rows(
        self,
        rows: list[dict[str, str]],
        headers: list[str],
    ) -> dict[str, Any]:
        brand_header = self._match_header(headers, _BRAND_HEADER_KEYWORDS)
        competitor_header = self._match_header(headers, _COMPETITOR_HEADER_KEYWORDS)
        industry_header = self._match_header(headers, _INDUSTRY_HEADER_KEYWORDS)
        keyword_header = self._match_header(headers, _KEYWORD_HEADER_KEYWORDS)
        product_header = self._match_header(headers, _PRODUCT_HEADER_KEYWORDS)
        website_header = self._match_header(headers, ("官网", "website", "url", "domain"))

        brand_profile_patch: dict[str, Any] = {}
        competitors: list[dict[str, Any]] = []
        brand_keywords: list[str] = []
        core_products: list[str] = []

        for row in rows:
            if brand_header and not brand_profile_patch.get("brand_name"):
                brand_name = str(row.get(brand_header, "")).strip()
                if brand_name:
                    brand_profile_patch["brand_name"] = brand_name
            if website_header and not brand_profile_patch.get("official_website"):
                website = str(row.get(website_header, "")).strip()
                if website:
                    brand_profile_patch["official_website"] = website
            if industry_header and not brand_profile_patch.get("industry"):
                industry = str(row.get(industry_header, "")).strip()
                if industry:
                    brand_profile_patch["industry"] = industry
            if keyword_header:
                brand_keywords.extend(self._split_multi_values(str(row.get(keyword_header, "")).strip()))
            if product_header:
                core_products.extend(self._split_multi_values(str(row.get(product_header, "")).strip()))
            if competitor_header:
                for competitor_name in self._split_multi_values(
                    str(row.get(competitor_header, "")).strip()
                ):
                    competitors.append({"name": competitor_name})

        if brand_keywords:
            brand_profile_patch["brand_keywords"] = list(dict.fromkeys(brand_keywords))
        if core_products:
            brand_profile_patch["core_products"] = list(dict.fromkeys(core_products))

        deduped_competitors = []
        seen = set()
        for competitor in competitors:
            name = str(competitor.get("name") or "").strip()
            if not name or name in seen:
                continue
            seen.add(name)
            deduped_competitors.append(competitor)

        return {
            "brand_profile_patch": brand_profile_patch,
            "competitors": deduped_competitors,
        }

    def _build_warnings(
        self, parsed: ParsedTable, *, skipped_question_count: int = 0
    ) -> list[str]:
        warnings: list[str] = []
        if parsed.empty_row_count:
            warnings.append(f"已跳过 {parsed.empty_row_count} 行空数据")
        if parsed.duplicate_row_count:
            warnings.append(f"已跳过 {parsed.duplicate_row_count} 行重复数据")
        if skipped_question_count:
            warnings.append(f"已跳过 {skipped_question_count} 行未识别到问题内容")
        return warnings

    def _match_header(self, headers: list[str], keywords: tuple[str, ...]) -> str | None:
        best_header: str | None = None
        best_score = 0
        for header in headers:
            score = self._header_match_score(header, keywords)
            if score > best_score:
                best_header = header
                best_score = score
        return best_header

    def _header_match_score(self, header: str, keywords: tuple[str, ...]) -> int:
        lowered = header.lower().strip()
        compact = lowered.replace(" ", "").replace("_", "")
        score = 0

        for keyword in keywords:
            keyword_lower = keyword.lower().strip()
            keyword_compact = keyword_lower.replace(" ", "").replace("_", "")
            if compact == keyword_compact:
                score = max(score, 100)
            elif compact.startswith(keyword_compact):
                score = max(score, 80)
            elif keyword_compact in compact:
                score = max(score, 60)

        if score == 0:
            return 0

        if any(marker in compact for marker in ("id", "编号", "序号", "序列", "编码")):
            score -= 35
        if any(marker in compact for marker in ("内容", "文本", "text", "detail", "详情", "提问")):
            score += 12

        return max(score, 0)

    def _contains_any(self, value: str, keywords: tuple[str, ...]) -> bool:
        lowered = value.lower()
        return any(keyword in lowered for keyword in keywords)

    def _looks_like_link_sheet(self, parsed: ParsedTable) -> bool:
        url_count = 0
        inspected = 0
        for row in parsed.rows[:10]:
            for value in row.values():
                text = str(value).strip()
                if not text:
                    continue
                inspected += 1
                if self._looks_like_url(text):
                    url_count += 1
        return inspected > 0 and url_count >= max(2, inspected // 3)

    def _looks_like_brand_sheet(self, parsed: ParsedTable) -> bool:
        score = 0
        for header in parsed.headers:
            if self._contains_any(header, _BRAND_HEADER_KEYWORDS):
                score += 2
            if self._contains_any(header, _COMPETITOR_HEADER_KEYWORDS):
                score += 2
        return score >= 2

    def _looks_like_freeform_question_rows(
        self,
        rows: list[dict[str, str]],
        only_header: str,
    ) -> bool:
        sample_values = [str(row.get(only_header, "")).strip() for row in rows[:8]]
        populated = [value for value in sample_values if value]
        if len(populated) < 2:
            return False
        question_like = sum(
            1
            for value in populated
            if any(token in value for token in ("?", "？", "怎么", "如何", "为什么", "推荐", "哪", "是否"))
            or len(value) >= 8
        )
        return question_like >= max(2, len(populated) // 2)

    def _looks_like_url(self, value: str) -> bool:
        lowered = value.lower()
        return lowered.startswith("http://") or lowered.startswith("https://") or lowered.startswith("www.")

    def _split_multi_values(self, value: str) -> list[str]:
        if not value:
            return []
        normalized = value
        for separator in ("；", ";", "，", ",", "|", "/", "、"):
            normalized = normalized.replace(separator, "\n")
        return [item.strip() for item in normalized.splitlines() if item.strip()]
